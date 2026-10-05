#!/usr/bin/env python3
"""
eu.py – Fetch EU clinical trials from both CTIS and EudraCT, deduplicate
        on trial ID, and write the results to an Excel workbook.

Usage:
    python eu.py Semaglutide
    python eu.py Semaglutide --no-details --max-records 50
    python eu.py Semaglutide --out my_results.xlsx
    python eu.py Semaglutide --workers 10 --eudract-delay 0.5

Both registers are queried concurrently, and per-trial detail pages within each
register are fetched by a thread pool (see --workers).  ctis.py / eudract.py are
used as libraries and are not modified.
"""

import argparse
import re
import sys
import threading
import time
from concurrent.futures import Future, ThreadPoolExecutor, as_completed
from typing import Any, Callable, Dict, List, Optional

import requests

try:
    import pandas as pd
except ImportError:
    pd = None

try:
    from fetch_search_terms import fetch_search_terms as _fetch_search_terms
except ImportError:
    _fetch_search_terms = None

# ── Local imports (ctis.py and eudract.py must be in the same directory) ──────
# The modules are imported (not just their public functions) because the
# parallel orchestrator below reuses their private search / detail / flatten
# helpers.  ctis.py and eudract.py themselves are left untouched.
try:
    import ctis as _ctis_mod
    from ctis import ctis
except ImportError:
    _ctis_mod = None  # type: ignore
    ctis = None  # type: ignore
    print("[eu] WARNING: ctis.py not found – CTIS results will be skipped.", file=sys.stderr)

try:
    import eudract as _eudract_mod
    from eudract import eudract
except ImportError:
    _eudract_mod = None  # type: ignore
    eudract = None  # type: ignore
    print("[eu] WARNING: eudract.py not found – EudraCT results will be skipped.", file=sys.stderr)

try:
    import openpyxl
    from openpyxl.styles import (
        Font, PatternFill, Alignment, Border, Side
    )
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl is required:  pip install openpyxl")


# ── Deduplication ─────────────────────────────────────────────────────────────

TRIAL_ID_COL = "Trial ID"
_ID_SEP      = "; "


def _row_ids(row: Dict[str, Any]) -> List[str]:
    """All identifiers a row carries (CTIS ct_number and/or EudraCT number)."""
    ids = []
    for field in ("ct_number", "eudract_number"):
        v = str(row.get(field) or "").strip()
        if v and v not in ids:
            ids.append(v)
    return ids


def _merge_rows(rows: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """
    Deduplicate rows on trial identifier and build a single "Trial ID" column.

    Two rows are the same trial if they share *any* identifier (CTIS ct_number
    or EudraCT number).  Links are followed transitively, so a CTIS row
    (ct_number + EudraCT number) bridges to its legacy EudraCT row, and
    duplicates within one source are collapsed too.

    The first row seen is kept; empty fields are back-filled from later
    duplicates.  The separate ct_number / eudract_number columns are replaced
    by "Trial ID", which holds every ID of the trial ("; "-separated).
    Output order = order of first appearance.
    """
    groups: List[Dict[str, Any]] = []        # merged rows, in first-seen order
    owner:  Dict[str, int]       = {}        # identifier -> index into groups
    redirect: Dict[int, int]     = {}         # group merged into another group

    def resolve(i: int) -> int:
        while i in redirect:
            i = redirect[i]
        return i

    for row in rows:
        ids = _row_ids(row)
        hits = sorted({resolve(owner[i]) for i in ids if i in owner})

        if not hits:                                   # brand-new trial
            groups.append(dict(row))
            gi = len(groups) - 1
        else:                                          # duplicate → fold in
            gi = hits[0]
            target = groups[gi]
            sources = {str(target.get("source", "")), str(row.get("source", ""))}
            for k, v in row.items():
                if not target.get(k) and v:
                    target[k] = v
            # a row bridging two previously separate groups: merge them too
            for other in hits[1:]:
                for k, v in groups[other].items():
                    if not target.get(k) and v:
                        target[k] = v
                sources.add(str(groups[other].get("source", "")))
                redirect[other] = gi
                groups[other]   = None                 # type: ignore
            sources.discard("")
            if len(sources) > 1:
                target["source"] = " + ".join(sorted(sources, key=str.upper))
        for i in ids:
            owner[i] = gi

    merged: List[Dict[str, Any]] = []
    for g in groups:
        if g is None:
            continue
        ids = _row_ids(g)
        out: Dict[str, Any] = {TRIAL_ID_COL: _ID_SEP.join(ids)}
        for k, v in g.items():
            if k not in ("ct_number", "eudract_number", "trial_id"):
                out[k] = v
        merged.append(out)
    return merged


# ── Excel output ──────────────────────────────────────────────────────────────

# Columns that always appear first (in this order) when present in the data
PRIORITY_COLUMNS = [
    "Trial ID",        # CTIS ct_number + EudraCT number, deduplicated
    "source",
    "title",
    "short_title",
    "status",
    "phase",
    "conditions",
    "sponsor",
    "sponsor_type",
    "countries",
    "start_date",
    "end_date",
    "enrolled",
    "age_group",
    "gender",
    "primary_endpoint",
    "secondary_endpoint",
    "decision_date",
    "last_updated",
    "results_available",
    "url",
]

# Colours
_HEADER_FILL   = PatternFill("solid", fgColor="1F4E79")   # dark blue
_ALT_FILL      = PatternFill("solid", fgColor="D6E4F0")   # light blue
_PLAIN_FILL    = PatternFill("solid", fgColor="FFFFFF")   # white
_HEADER_FONT   = Font(name="Arial", bold=True, color="FFFFFF", size=10)
_BODY_FONT     = Font(name="Arial", size=10)
_BORDER_SIDE   = Side(style="thin", color="BDC3C7")
_CELL_BORDER   = Border(
    left=_BORDER_SIDE, right=_BORDER_SIDE,
    top=_BORDER_SIDE,  bottom=_BORDER_SIDE,
)
_WRAP = Alignment(wrap_text=True, vertical="top")


def _ordered_columns(rows: List[Dict[str, Any]]) -> List[str]:
    all_keys: List[str] = []
    for row in rows:
        for k in row:
            if k not in all_keys:
                all_keys.append(k)
    ordered = [c for c in PRIORITY_COLUMNS if c in all_keys]
    remainder = [c for c in all_keys if c not in ordered]
    return ordered + remainder


def write_excel(rows: List[Dict[str, Any]], path: str, drug_name: str):
    """Write *rows* to an Excel workbook at *path*."""
    if not rows:
        print("[eu] No rows to write.", file=sys.stderr)
        return

    wb     = openpyxl.Workbook()
    ws     = wb.active
    ws.title = "EU Clinical Trials"

    columns = _ordered_columns(rows)

    # ── Header row ────────────────────────────────────────────────────────────
    for col_idx, col_name in enumerate(columns, start=1):
        cell            = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font       = _HEADER_FONT
        cell.fill       = _HEADER_FILL
        cell.alignment  = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border     = _CELL_BORDER

    ws.row_dimensions[1].height = 30

    # ── Data rows ─────────────────────────────────────────────────────────────
    for row_idx, row in enumerate(rows, start=2):
        fill = _ALT_FILL if row_idx % 2 == 0 else _PLAIN_FILL
        for col_idx, col_name in enumerate(columns, start=1):
            val  = row.get(col_name, "")
            # Convert lists/dicts to string for readability
            if isinstance(val, (list, dict)):
                val = str(val)
            cell           = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font      = _BODY_FONT
            cell.fill      = fill
            cell.alignment = _WRAP
            cell.border    = _CELL_BORDER

    # ── Column widths (auto-size, capped) ────────────────────────────────────
    WIDE_COLS   = {"title", "short_title", "conditions", "primary_endpoint",
                   "secondary_endpoint", "sponsor", "inclusion_criteria",
                   "exclusion_criteria", "findings", "main_objective"}
    NARROW_COLS = {"source", "phase", "gender", "age_group", "enrolled"}

    for col_idx, col_name in enumerate(columns, start=1):
        col_letter = get_column_letter(col_idx)
        if col_name in WIDE_COLS:
            ws.column_dimensions[col_letter].width = 45
        elif col_name == "url":
            ws.column_dimensions[col_letter].width = 50
        elif col_name == TRIAL_ID_COL:
            ws.column_dimensions[col_letter].width = 36
        elif col_name in NARROW_COLS:
            ws.column_dimensions[col_letter].width = 18
        else:
            ws.column_dimensions[col_letter].width = 25

    # ── Freeze header & auto-filter ───────────────────────────────────────────
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # ── Summary sheet ─────────────────────────────────────────────────────────
    ws2        = wb.create_sheet("Summary")
    ws2.title  = "Summary"

    ctis_rows    = [r for r in rows if str(r.get("source", "")).upper() == "CTIS"]
    eudract_rows = [r for r in rows if str(r.get("source", "")).upper() == "EUDRACT"]
    both_rows    = [r for r in rows if str(r.get("source", "")).upper() not in ("CTIS", "EUDRACT")]

    summary_data = [
        ["EU Clinical Trials – Summary"],
        [],
        ["Drug / Search term", drug_name],
        ["Total unique trials",  len(rows)],
        ["  from CTIS",          len(ctis_rows)],
        ["  from EudraCT",       len(eudract_rows)],
        ["  merged / other",     len(both_rows)],
        [],
        ["Status breakdown"],
    ]
    status_counts: Dict[str, int] = {}
    for r in rows:
        s = str(r.get("status", "Unknown") or "Unknown")
        status_counts[s] = status_counts.get(s, 0) + 1
    for status, count in sorted(status_counts.items(), key=lambda x: -x[1]):
        summary_data.append([f"  {status}", count])

    summary_data += [
        [],
        ["Phase breakdown"],
    ]
    phase_counts: Dict[str, int] = {}
    for r in rows:
        p = str(r.get("phase", "Unknown") or "Unknown")
        phase_counts[p] = phase_counts.get(p, 0) + 1
    for phase, count in sorted(phase_counts.items(), key=lambda x: -x[1]):
        summary_data.append([f"  {phase}", count])

    for r_idx, data_row in enumerate(summary_data, start=1):
        for c_idx, val in enumerate(data_row, start=1):
            cell       = ws2.cell(row=r_idx, column=c_idx, value=val)
            cell.font  = Font(name="Arial", size=10,
                              bold=(r_idx == 1 or val in ("Status breakdown", "Phase breakdown")))
    ws2.column_dimensions["A"].width = 30
    ws2.column_dimensions["B"].width = 15

    wb.save(path)
    print(f"[eu] Wrote {len(rows)} unique trial(s) → {path}", file=sys.stderr)


# ── Parallel fetch infrastructure ─────────────────────────────────────────────

DEFAULT_WORKERS        = 6      # concurrent detail fetches per source
DEFAULT_CTIS_DELAY     = 0.10   # min seconds between CTIS detail request starts
DEFAULT_EUDRACT_DELAY  = 0.30   # min seconds between EudraCT detail request starts

_print_lock = threading.Lock()


def _log(msg: str) -> None:
    """Thread-safe stderr print so lines from parallel workers don't interleave."""
    with _print_lock:
        print(msg, file=sys.stderr)


class _RateLimiter:
    """
    Global (per-source) throttle shared by all workers of that source.

    Guarantees that request *start times* are at least ``min_interval`` seconds
    apart, while still letting the (slow) network round-trips overlap.  This
    replaces the per-request ``time.sleep`` calls of the sequential modules.
    """

    def __init__(self, min_interval: float):
        self._interval = max(0.0, min_interval)
        self._lock     = threading.Lock()
        self._next     = 0.0

    def wait(self) -> None:
        if self._interval <= 0:
            return
        with self._lock:
            now         = time.monotonic()
            slot        = max(now, self._next)
            self._next  = slot + self._interval
        delay = slot - now
        if delay > 0:
            time.sleep(delay)


class _SessionPool:
    """One ``requests.Session`` per worker thread (Session isn't guaranteed thread-safe)."""

    def __init__(self, headers: Dict[str, str], pool_size: int):
        self._headers = headers
        self._pool    = pool_size
        self._local   = threading.local()

    def get(self) -> requests.Session:
        s = getattr(self._local, "session", None)
        if s is None:
            s = requests.Session()
            s.headers.update(self._headers)
            adapter = requests.adapters.HTTPAdapter(pool_connections=self._pool,
                                                    pool_maxsize=self._pool)
            s.mount("https://", adapter)
            s.mount("http://",  adapter)
            self._local.session = s
        return s


def _run_pipeline(label:       str,
                  records:     Any,                       # iterable of search-result records
                  work:        Callable[[Any], Any],      # record -> finished row
                  fallback:    Callable[[Any], Any],      # record -> row if work() raised
                  workers:     int,
                  progress_every: int,
                  verbose:     bool) -> List[Any]:
    """
    Consume the (sequential, paginated) search generator in the calling thread
    and hand every record to a worker pool as soon as it arrives, so detail
    fetches overlap with the fetching of later search pages.

    Output order == search order, regardless of which worker finishes first.
    If the search itself fails part-way, the rows gathered so far are kept.
    """
    done_count = 0
    count_lock = threading.Lock()

    def _on_done(_fut: Future) -> None:
        nonlocal done_count
        with count_lock:
            done_count += 1
            n = done_count
        if verbose and n % progress_every == 0:
            _log(f"  … {n} {label} trials processed")

    pool    = ThreadPoolExecutor(max_workers=max(1, workers),
                                 thread_name_prefix=f"eu-{label.lower()}")
    pending: List[tuple] = []     # (record, future) in search order
    try:
        try:
            for rec in records:
                fut = pool.submit(work, rec)
                fut.add_done_callback(_on_done)
                pending.append((rec, fut))
        except Exception as exc:
            _log(f"[eu] {label} search aborted after {len(pending)} record(s): {exc}"
                 f" – keeping what was fetched so far.")

        rows: List[Any] = []
        for rec, fut in pending:
            try:
                rows.append(fut.result())
            except Exception as exc:
                _log(f"  ! {label} worker failed ({exc}); using search-level data only.")
                try:
                    rows.append(fallback(rec))
                except Exception:
                    pass
        return rows
    except BaseException:
        # Ctrl-C etc.: don't leave queued work running
        pool.shutdown(wait=False, cancel_futures=True)
        raise
    finally:
        pool.shutdown(wait=True)


# ── CTIS (parallel detail fetch) ──────────────────────────────────────────────

def fetch_ctis_parallel(drug_name:   str,
                        details:     bool,
                        max_records: Optional[int],
                        workers:     int,
                        delay:       float,
                        verbose:     bool) -> List[Dict[str, Any]]:
    m        = _ctis_mod
    sessions = _SessionPool(m.HEADERS, workers + 1)
    limiter  = _RateLimiter(delay)

    if verbose:
        _log(f"[CTIS] Searching for '{drug_name}' ({workers} worker(s)) …")

    search = m._search_trials(drug_name, sessions.get(),
                              page_size=50, max_records=max_records, verbose=verbose)

    def work(summary: Dict[str, Any]) -> Dict[str, Any]:
        det = None
        if details:
            ct = summary.get("ctNumber", "")
            limiter.wait()
            try:
                det = m._get_trial_details(ct, sessions.get())
            except Exception as exc:
                if verbose:
                    _log(f"  ! detail fetch failed for {ct}: {exc}")
        return m._flatten(summary, det)

    rows = _run_pipeline("CTIS", search, work,
                         fallback=lambda s: m._flatten(s, None),
                         workers=workers, progress_every=25, verbose=verbose)
    if verbose:
        _log(f"[CTIS] Done – {len(rows)} trial(s) fetched.")
    return rows


# ── EudraCT (parallel detail + results fetch) ─────────────────────────────────

def fetch_eudract_parallel(drug_name:   str,
                           details:     bool,
                           max_records: Optional[int],
                           workers:     int,
                           delay:       float,
                           verbose:     bool) -> List[Dict[str, Any]]:
    m        = _eudract_mod
    sessions = _SessionPool(m.HEADERS, workers + 1)
    limiter  = _RateLimiter(delay)

    if verbose:
        _log(f"[EudraCT] Searching for '{drug_name}' ({workers} worker(s)) …")

    search = m._search_trials(drug_name, sessions.get(),
                              max_records=max_records, verbose=verbose)

    def work(row: Dict[str, Any]) -> Dict[str, Any]:
        # Each row is owned by exactly one worker, so in-place update is safe.
        sess = sessions.get()
        if details:
            country = (row.get("countries", "").split(";")[0] or "GB").strip()
            limiter.wait()
            try:
                row.update(m._get_trial_details(row["eudract_number"], country, sess))
            except Exception as exc:
                if verbose:
                    _log(f"  ! detail fetch failed for {row['eudract_number']}: {exc}")
        # Results are fetched regardless of `details` (same as eudract() default).
        if row.get("results_available") == "Yes":
            limiter.wait()
            try:
                row.update(m._get_trial_results(row["eudract_number"], sess))
            except Exception as exc:
                if verbose:
                    _log(f"  ! results fetch failed for {row['eudract_number']}: {exc}")
        return row

    rows = _run_pipeline("EudraCT", search, work,
                         fallback=lambda r: r,
                         workers=workers, progress_every=20, verbose=verbose)
    if verbose:
        _log(f"[EudraCT] Done – {len(rows)} trial(s) fetched.")
    return rows


# ── Orchestrator ──────────────────────────────────────────────────────────────

def fetch_eu(drug_name:     str,
             details:       bool          = True,
             max_records:   Optional[int] = None,
             verbose:       bool          = True,
             workers:       int           = DEFAULT_WORKERS,
             ctis_delay:    float         = DEFAULT_CTIS_DELAY,
             eudract_delay: float         = DEFAULT_EUDRACT_DELAY,
             ) -> List[Dict[str, Any]]:
    """
    Fetch trials from CTIS + EudraCT **in parallel**, merge duplicates, return
    unique rows.

    Parallelism happens on two levels:
      1. CTIS and EudraCT are queried concurrently.
      2. Within each source, per-trial detail/results pages are fetched by
         *workers* threads (throttled by a per-source minimum request interval).

    Row order is deterministic: all CTIS rows (search order) followed by all
    EudraCT rows (search order), exactly as in the sequential version.
    """
    jobs: Dict[str, Future] = {}
    outer = ThreadPoolExecutor(max_workers=2, thread_name_prefix="eu-source")
    try:
        if _ctis_mod is not None:
            jobs["CTIS"] = outer.submit(fetch_ctis_parallel, drug_name, details,
                                        max_records, workers, ctis_delay, verbose)
        else:
            _log("[eu] ctis module unavailable – skipping.")

        if _eudract_mod is not None:
            jobs["EudraCT"] = outer.submit(fetch_eudract_parallel, drug_name, details,
                                           max_records, workers, eudract_delay, verbose)
        else:
            _log("[eu] eudract module unavailable – skipping.")

        all_rows: List[Dict[str, Any]] = []
        for name in ("CTIS", "EudraCT"):          # fixed order → deterministic merge
            fut = jobs.get(name)
            if fut is None:
                continue
            try:
                src_rows = fut.result()
                all_rows.extend(src_rows)
                if verbose:
                    _log(f"[eu] {name} returned {len(src_rows)} trial(s).")
            except Exception as exc:
                _log(f"[eu] {name} fetch error: {exc}")
    except BaseException:
        outer.shutdown(wait=False, cancel_futures=True)
        raise
    finally:
        outer.shutdown(wait=True)

    # ── Deduplicate ───────────────────────────────────────────────────────────
    merged = _merge_rows(all_rows)
    if verbose:
        _log(f"[eu] {len(all_rows)} total rows → {len(merged)} unique trial(s) after merge.")
    return merged


# ── Importable entry point ────────────────────────────────────────────────────

def _map_source(raw: str):
    """Map eu.py's internal 'source' field to (Source, Registry) for the standard output."""
    low = raw.lower().strip()
    if "ctis" in low and "eudract" in low:
        return "CTIS, EudraCT", "EU CTIS / EudraCT"
    if "ctis" in low:
        return "CTIS", "EU CTIS"
    if "eudract" in low:
        return "EudraCT", "EudraCT"
    return raw, ""


def eu(
    drug_name: str,
    fetch_terms: bool = True,
    details: bool = True,
    workers: int = DEFAULT_WORKERS,
    ctis_delay: float = DEFAULT_CTIS_DELAY,
    eudract_delay: float = DEFAULT_EUDRACT_DELAY,
    term_workers: int = 6,
) -> "pd.DataFrame":
    """
    Main entry point — importable.

        from eu import eu
        df = eu("Semaglutide")   # -> DataFrame with Trial_ID | Registry | Source

    Searches CTIS and EudraCT for every synonym returned by fetch_search_terms.
    Source is "CTIS", "EudraCT", or "CTIS, EudraCT" for trials found in both registers.
    """
    if pd is None:
        raise ImportError("pip install pandas")

    drug = " ".join(str(drug_name or "").split())
    if not drug:
        raise ValueError("drug_name is required")

    # 1. Build search-term list
    terms = [drug]
    if fetch_terms and _fetch_search_terms is not None:
        try:
            fetched = _fetch_search_terms(drug)
            seen_terms = {drug.lower()}
            for t in fetched:
                if t.lower() not in seen_terms:
                    seen_terms.add(t.lower())
                    terms.append(t)
        except Exception as e:
            print(f"  [eu] term lookup failed ({e}) — searching only '{drug}'")
    print(f"  [eu] search terms: {terms}")

    # 2. Fetch for each term in parallel; deduplicate across calls on individual IDs.
    #    A merged row may hold multiple IDs in TRIAL_ID_COL ("id1; id2"), so we
    #    track every individual component to avoid repeating the same trial.
    all_rows: list[dict] = []
    seen_ids: set[str] = set()
    _lock = threading.Lock()

    def _fetch_term(term: str) -> list[dict]:
        try:
            return fetch_eu(
                term,
                details=details,
                workers=workers,
                ctis_delay=ctis_delay,
                eudract_delay=eudract_delay,
                verbose=False,
            )
        except Exception as e:
            print(f"  [eu] fetch failed for term '{term}': {e}")
            return []

    with ThreadPoolExecutor(max_workers=min(term_workers, len(terms)),
                            thread_name_prefix="eu-term") as pool:
        futures = {pool.submit(_fetch_term, t): t for t in terms}
        for fut in as_completed(futures):
            term_rows = fut.result()
            with _lock:
                for row in term_rows:
                    trial_id_str = str(row.get(TRIAL_ID_COL, "")).strip()
                    if not trial_id_str:
                        continue
                    # Individual IDs within a compound "id1; id2" value
                    parts = [p.strip() for p in trial_id_str.split(_ID_SEP) if p.strip()]
                    if any(p in seen_ids for p in parts):
                        continue  # already captured from another term
                    seen_ids.update(parts)
                    all_rows.append(row)

    if not all_rows:
        print("  [eu] no trials found.")
        return pd.DataFrame(columns=["Trial_ID", "Registry", "Source"])

    # 3. Build standard DataFrame
    result_rows = []
    for row in all_rows:
        trial_id = str(row.get(TRIAL_ID_COL, "")).strip()
        source, registry = _map_source(str(row.get("source", "")))
        result_rows.append({"Trial_ID": trial_id, "Registry": registry, "Source": source})

    result = pd.DataFrame(result_rows)
    print(f"  [eu] {len(result)} unique trials found across {len(terms)} search term(s).")
    return result


# ── CLI ───────────────────────────────────────────────────────────────────────

def main():
    ap = argparse.ArgumentParser(
        description=(
            "Fetch EU clinical trials from CTIS + EudraCT and write unique "
            "trial IDs to an Excel workbook.\n\n"
            "Example:  python eu.py Semaglutide"
        ),
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    ap.add_argument("drug", help="Drug / INN / brand name to search for")
    ap.add_argument("--no-details", dest="details", action="store_false",
                    help="Skip per-trial detail page fetch (faster)")
    ap.add_argument("--max-records", type=int, default=None,
                    help="Limit results per source (default: all)")
    ap.add_argument("--out", default=None,
                    help="Output .xlsx path (default: <drug>_eu_trials.xlsx)")
    ap.add_argument("--quiet", action="store_true",
                    help="Suppress progress messages")
    ap.add_argument("--workers", type=int, default=DEFAULT_WORKERS,
                    help=f"Parallel detail-fetch threads per source "
                         f"(default: {DEFAULT_WORKERS}; 1 = effectively sequential)")
    ap.add_argument("--ctis-delay", type=float, default=DEFAULT_CTIS_DELAY,
                    help=f"Min seconds between CTIS detail requests "
                         f"(default: {DEFAULT_CTIS_DELAY})")
    ap.add_argument("--eudract-delay", type=float, default=DEFAULT_EUDRACT_DELAY,
                    help=f"Min seconds between EudraCT detail requests "
                         f"(default: {DEFAULT_EUDRACT_DELAY})")
    args = ap.parse_args()
    if args.workers < 1:
        ap.error("--workers must be >= 1")

    out_path = args.out or f"{args.drug.lower().replace(' ', '_')}_eu_trials.xlsx"
    verbose  = not args.quiet

    rows = fetch_eu(args.drug, details=args.details,
                    max_records=args.max_records, verbose=verbose,
                    workers=args.workers,
                    ctis_delay=args.ctis_delay,
                    eudract_delay=args.eudract_delay)

    if not rows:
        print("[eu] No trials found.", file=sys.stderr)
        return 0

    write_excel(rows, out_path, args.drug)

    # Brief console preview
    print(f"\n{'─'*90}", file=sys.stderr)
    print(f"  {'Trial ID':<36}  {'Phase':<10}  {'Status':<20}  Title", file=sys.stderr)
    print(f"{'─'*90}", file=sys.stderr)
    for row in rows[:10]:
        tid    = str(row.get(TRIAL_ID_COL, ""))
        phase  = str(row.get("phase", ""))[:8]
        status = str(row.get("status", ""))[:18]
        title  = str(row.get("title", ""))[:40]
        print(f"  {tid:<36}  {phase:<10}  {status:<20}  {title}", file=sys.stderr)
    if len(rows) > 10:
        print(f"  … and {len(rows) - 10} more.", file=sys.stderr)
    print(f"{'─'*90}", file=sys.stderr)

    return 0


if __name__ == "__main__":
    sys.exit(main())